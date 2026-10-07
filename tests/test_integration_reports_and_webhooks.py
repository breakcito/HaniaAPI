"""Pruebas de integración de Reportería en Excel y Recepción de Webhooks."""

import pytest
from io import BytesIO
import openpyxl
from fastapi.testclient import TestClient

def test_sales_excel_report_endpoint(client: TestClient, auth_headers: dict, test_company):
    """Verifica la generación y descarga de reporte de ventas en formato binario XLSX."""
    resp = client.get(
        f"/api/reports/sales-excel?company_id={test_company.id}",
        headers=auth_headers
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=" in resp.headers.get("content-disposition", "")

    # Validar que sea un archivo Excel válido legible por openpyxl
    content = resp.content
    assert len(content) > 1000
    wb = openpyxl.load_workbook(BytesIO(content))
    assert len(wb.sheetnames) >= 1

def test_despatches_excel_report_endpoint(client: TestClient, auth_headers: dict, test_company):
    """Verifica la generación y descarga de reporte de Guías de Remisión en Excel."""
    resp = client.get(
        f"/api/reports/despatches-excel?company_id={test_company.id}",
        headers=auth_headers
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    content = resp.content
    assert len(content) > 1000
    wb = openpyxl.load_workbook(BytesIO(content))
    assert len(wb.sheetnames) >= 1

def test_webhook_receiver_event(client: TestClient):
    """Verifica el receptor de eventos asíncronos de Factos API."""
    webhook_payload = {
        "event": "document.accepted",
        "data": {
            "id": 9999,
            "type_code": "01",
            "series": "F001",
            "correlative": 88,
            "status": "accepted",
            "pdf_url": "https://api-factos.test/pdf/9999",
            "xml_url": "https://api-factos.test/xml/9999",
            "cdr_url": "https://api-factos.test/cdr/9999",
            "sunat_description": "La Factura F001-00000088 ha sido aceptada."
        }
    }
    resp = client.post("/api/webhooks/factos", json=webhook_payload)
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "success"
    assert data["event"] == "document.accepted"
